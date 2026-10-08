"""
DocuMind Document Loader.
Extracts structured text units with source, page, and location metadata
from PDF, DOCX, PPTX, TXT, and legacy DOC/PPT files.
"""

import logging
import shutil
import subprocess
from pathlib import Path
from typing import Any, Dict, List, Optional

logger = logging.getLogger(__name__)

# PDF parser (prefer modern pymupdf over deprecated fitz name)
try:
    import pymupdf as fitz
except ImportError:
    try:
        import fitz
    except ImportError:
        fitz = None

# DOCX parser
try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

# PPTX parser
try:
    from pptx import Presentation
except ImportError:
    Presentation = None


def _find_libreoffice() -> Optional[str]:
    """Locate LibreOffice / soffice binary across platforms for legacy conversions."""
    candidates = [
        "libreoffice",
        "soffice",
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]
    for cmd in candidates:
        if shutil.which(cmd):
            return cmd
    return None


def convert_legacy(path: Path, target_ext: str) -> Path:
    """
    Convert legacy binary formats (.doc -> .docx, .ppt -> .pptx) via headless LibreOffice.
    """
    lo = _find_libreoffice()
    if not lo:
        raise RuntimeError(
            "Headless LibreOffice is required to convert legacy .doc/.ppt files. "
            "Please upload modern formats (.docx/.pptx) or install LibreOffice."
        )

    out_dir = path.parent
    try:
        result = subprocess.run(
            [
                lo,
                "--headless",
                "--convert-to",
                target_ext,
                "--outdir",
                str(out_dir),
                str(path),
            ],
            capture_output=True,
            text=True,
            timeout=120,
        )
    except subprocess.TimeoutExpired:
        raise RuntimeError(f"Legacy document conversion timed out for {path.name}.")

    if result.returncode != 0:
        logger.error("LibreOffice conversion failed: %s", result.stderr)
        raise RuntimeError(f"Legacy conversion failed: {result.stderr[:200]}")

    converted = out_dir / (path.stem + "." + target_ext)
    if not converted.exists():
        raise RuntimeError("Converted output file was not generated.")
    return converted


def load_pdf(path: Path) -> List[Dict[str, Any]]:
    """Extract page-indexed text units from a PDF file."""
    if fitz is None:
        raise RuntimeError("PyMuPDF (pymupdf) is not installed.")

    try:
        doc = fitz.open(str(path))
    except Exception as exc:
        logger.error("Failed to open PDF %s: %s", path.name, exc)
        raise ValueError(f"Corrupted or unreadable PDF file: {exc}")

    units: List[Dict[str, Any]] = []
    try:
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if text and text.strip():
                units.append(
                    {
                        "text": text.strip(),
                        "source": path.name,
                        "location": f"Page {page_num}",
                        "page": page_num,
                        "type": "pdf",
                    }
                )
    finally:
        doc.close()

    return units


def load_docx(path: Path) -> List[Dict[str, Any]]:
    """Extract paragraph and table text from a Word document (.docx)."""
    if DocxDocument is None:
        raise RuntimeError("python-docx is not installed.")

    try:
        doc = DocxDocument(str(path))
    except Exception as exc:
        logger.error("Failed to open DOCX %s: %s", path.name, exc)
        raise ValueError(f"Corrupted or invalid DOCX file: {exc}")

    parts: List[str] = []
    for para in doc.paragraphs:
        if para.text and para.text.strip():
            parts.append(para.text.strip())

    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()
            )
            if row_text:
                parts.append(row_text)

    if not parts:
        return []

    full_text = "\n\n".join(parts)
    return [
        {
            "text": full_text,
            "source": path.name,
            "location": "Document",
            "page": None,
            "type": "docx",
        }
    ]


def load_pptx(path: Path) -> List[Dict[str, Any]]:
    """Extract slide-indexed text units, tables, and notes from a PowerPoint presentation (.pptx)."""
    if Presentation is None:
        raise RuntimeError("python-pptx is not installed.")

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        logger.error("Failed to open PPTX %s: %s", path.name, exc)
        raise ValueError(f"Corrupted or invalid PPTX file: {exc}")

    units: List[Dict[str, Any]] = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        texts: List[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    if para.text and para.text.strip():
                        texts.append(para.text.strip())
            if shape.has_table:
                for row in shape.table.rows:
                    row_text = " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                        if cell.text and cell.text.strip()
                    )
                    if row_text:
                        texts.append(row_text)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes and notes.strip():
                texts.append(f"[Speaker Notes] {notes.strip()}")

        if texts:
            units.append(
                {
                    "text": "\n".join(texts),
                    "source": path.name,
                    "location": f"Slide {slide_num}",
                    "page": slide_num,
                    "type": "pptx",
                }
            )
    return units


def load_txt(path: Path) -> List[Dict[str, Any]]:
    """Extract text from plain text files with automatic encoding fallback."""
    for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
        try:
            text = path.read_text(encoding=encoding)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    else:
        raise ValueError(f"Unable to decode text file '{path.name}' with supported encodings.")

    if not text.strip():
        return []

    return [
        {
            "text": text.strip(),
            "source": path.name,
            "location": "Document",
            "page": None,
            "type": "txt",
        }
    ]


def load_document(path: Path) -> List[Dict[str, Any]]:
    """
    Main document parsing entry point. Dispatches to specific loader by extension.
    Validates file existence, non-empty content, and handles parsing failures.
    """
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    if path.stat().st_size == 0:
        raise ValueError(f"File '{path.name}' is empty (0 bytes).")

    ext = path.suffix.lower()
    supported = {".pdf", ".doc", ".docx", ".ppt", ".pptx", ".txt"}
    if ext not in supported:
        raise ValueError(
            f"Unsupported file type '{ext}'. Supported formats: {', '.join(sorted(supported))}"
        )

    logger.info("Loading document: %s (type: %s, size: %d bytes)", path.name, ext, path.stat().st_size)

    if ext == ".pdf":
        return load_pdf(path)
    if ext == ".docx":
        return load_docx(path)
    if ext == ".pptx":
        return load_pptx(path)
    if ext == ".txt":
        return load_txt(path)

    # Legacy formats
    if ext == ".doc":
        converted = convert_legacy(path, "docx")
        units = load_docx(converted)
        for u in units:
            u["source"] = path.name
            u["type"] = "doc"
        return units

    if ext == ".ppt":
        converted = convert_legacy(path, "pptx")
        units = load_pptx(converted)
        for u in units:
            u["source"] = path.name
            u["type"] = "ppt"
        return units

    return []